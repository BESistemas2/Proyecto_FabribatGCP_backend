import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generar_excel_informe_conciliacion(acta_data: dict) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Informe Conciliación"
    ws.views.sheetView[0].showGridLines = True

    # --- ESTILOS ---
    font_titulo = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    font_subtitulo = Font(name="Calibri", size=10, italic=True, color="595959")
    font_seccion = Font(name="Calibri", size=11, bold=True, color="1F4E79")
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_normal = Font(name="Calibri", size=10)
    font_header_tabla = Font(name="Calibri", size=10, bold=True, color="FFFFFF")

    fill_navy = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    fill_subseccion = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    border_thin = Border(
        top=Side(style='thin', color='BFBFBF'),
        bottom=Side(style='thin', color='BFBFBF'),
        left=Side(style='thin', color='BFBFBF'),
        right=Side(style='thin', color='BFBFBF')
    )
    border_total = Border(
        top=Side(style='thin', color='000000'),
        bottom=Side(style='double', color='000000')
    )

    fmt_moneda = '"$"#,##0.00;[Red]("$"#,##0.00);"-"'

    # --- 1. ENCABEZADO PRINCIPAL ---
    ws['A1'] = "INFORME DE CONCILIACIÓN BANCARIA"
    ws['A1'].font = font_titulo

    ws['A2'] = f"Empresa: {acta_data.get('empresa')} | Periodo: {acta_data.get('periodo_nombre')}"
    ws['A2'].font = font_subtitulo

    # --- 2. METADATOS ---
    metadata = [
        ("Empresa:", acta_data.get('empresa')),
        ("Banco:", acta_data.get('banco')),
        ("Número de Cuenta:", acta_data.get('numero_cuenta')),
        ("Periodo de Conciliación:", f"{acta_data.get('fecha_inicio')} al {acta_data.get('fecha_fin')}"),
        ("Moneda:", "USD ($)")
    ]

    row_idx = 4
    for label, val in metadata:
        ws.cell(row=row_idx, column=1, value=label).font = font_bold
        ws.cell(row=row_idx, column=2, value=val).font = font_normal
        row_idx += 1

    row_idx += 1

    # --- 3. TABLA CONCILIACIÓN DE SALDOS ---
    ws.cell(row=row_idx, column=1, value="CONCILIACIÓN BANCARIA DE SALDOS").font = font_seccion
    row_idx += 1

    ws.cell(row=row_idx, column=1, value="Concepto").fill = fill_navy
    ws.cell(row=row_idx, column=1).font = font_header_tabla
    ws.cell(row=row_idx, column=2, value="Monto ($)").fill = fill_navy
    ws.cell(row=row_idx, column=2).font = font_header_tabla
    ws.cell(row=row_idx, column=2).alignment = Alignment(horizontal="right")
    row_idx += 1

    saldos_filas = [
        (f"Saldo según extracto bancario al {acta_data.get('fecha_inicio')}", acta_data.get('saldo_extracto_inicial', 0), False, False),
        ("(+) Créditos", acta_data.get('total_creditos', 0), False, False),
        ("(-) Débitos", acta_data.get('total_debitos', 0), False, False),
        ("Saldo Según Extracto Bancario (Calculado)", acta_data.get('saldo_extracto_calculado', 0), True, True),
        (f"Saldo contable en libros al {acta_data.get('fecha_fin')}", acta_data.get('saldo_libros', 0), False, False),
        ("Diferencia / Variación Unexplained", acta_data.get('diferencia', 0), True, True),
    ]

    for concepto, monto, is_bold, has_total_border in saldos_filas:
        c1 = ws.cell(row=row_idx, column=1, value=concepto)
        c2 = ws.cell(row=row_idx, column=2, value=monto)

        c1.font = font_bold if is_bold else font_normal
        c2.font = font_bold if is_bold else font_normal
        c2.number_format = fmt_moneda
        c2.alignment = Alignment(horizontal="right")

        if has_total_border:
            c1.border = border_total
            c2.border = border_total

        row_idx += 1

    row_idx += 2

    # --- 4. DETALLE DE PARTIDAS CONCILIATORIAS ---
    ws.cell(row=row_idx, column=1, value="DETALLE DE PARTIDAS CONCILIATORIAS").font = font_seccion
    row_idx += 1
    ws.cell(row=row_idx, column=1, value=f"Partidas pendientes y ajustes del mes de {acta_data.get('periodo_nombre')}").font = font_subtitulo
    row_idx += 2

    def render_tabla_partidas(titulo_partida: str, headers: list, items: list):
        nonlocal row_idx
        ws.cell(row=row_idx, column=1, value=titulo_partida).font = font_bold
        row_idx += 1

        # Headers tabla
        for col_i, h in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_i, value=h)
            cell.fill = fill_navy
            cell.font = font_header_tabla
            if h == "Monto ($)":
                cell.alignment = Alignment(horizontal="right")
        row_idx += 1

        total_partida = 0.0
        for item in items:
            ws.cell(row=row_idx, column=1, value=item.get('fecha')).font = font_normal
            ws.cell(row=row_idx, column=2, value=item.get('numero')).font = font_normal
            ws.cell(row=row_idx, column=3, value=item.get('concepto')).font = font_normal
            
            c_monto = ws.cell(row=row_idx, column=4, value=item.get('monto', 0))
            c_monto.font = font_normal
            c_monto.number_format = fmt_moneda
            c_monto.alignment = Alignment(horizontal="right")

            total_partida += float(item.get('monto', 0))
            row_idx += 1

        # Fila Total
        ws.cell(row=row_idx, column=1, value="Total").font = font_bold
        c_tot = ws.cell(row=row_idx, column=4, value=total_partida)
        c_tot.font = font_bold
        c_tot.number_format = fmt_moneda
        c_tot.alignment = Alignment(horizontal="right")

        ws.cell(row=row_idx, column=1).border = border_total
        ws.cell(row=row_idx, column=4).border = border_total
        row_idx += 2

    headers_p = ["Fecha", "N° Documento", "Concepto", "Monto ($)"]

    render_tabla_partidas(
        "B. Transacciones no conciliadas (Depósitos en Tránsito)",
        headers_p,
        acta_data.get('partidas_transito_banco', [])
    )

    render_tabla_partidas(
        "C. Notas de Débito (Cargos Bancarios No Registrados)",
        headers_p,
        acta_data.get('notas_debito', [])
    )

    render_tabla_partidas(
        "D. Notas de Crédito (Abonos Bancarios No Registrados)",
        headers_p,
        acta_data.get('notas_credito', [])
    )

    # Ajustar anchos de columnas
    ws.column_dimensions['A'].width = 45
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 45
    ws.column_dimensions['D'].width = 20

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output