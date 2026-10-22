import io
from decimal import Decimal
from typing import Any, Iterable, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FMT_MONEDA = '"$"#,##0.00;[Red]-"$"#,##0.00;"-"'

_F_TITULO = Font(name="Calibri", size=14, bold=True, color="1F4E79")
_F_SUB = Font(name="Calibri", size=10, italic=True, color="595959")
_F_SECCION = Font(name="Calibri", size=11, bold=True, color="1F4E79")
_F_BOLD = Font(name="Calibri", size=10, bold=True)
_F_NORMAL = Font(name="Calibri", size=10)
_F_HEADER = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
_F_ALERTA = Font(name="Calibri", size=10, bold=True, color="C00000")
_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
_BORDE_TOTAL = Border(top=Side(style="thin", color="000000"), bottom=Side(style="double", color="000000"))
_DERECHA = Alignment(horizontal="right")


def _texto_seguro(v: Any) -> Any:
    """Evita que un texto del banco o de BC se interprete como fórmula de Excel."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    return v


def _num(v: Any):
    return float(v) if isinstance(v, Decimal) else v


def _suma(items: Iterable[dict]) -> Decimal:
    return sum((Decimal(str(i.get("monto", 0))) for i in items), Decimal("0"))


def generar_excel_informe_conciliacion(d: dict) -> io.BytesIO:
    """Informe de conciliación bancaria (ver ConciliacionService.obtener_datos_reporte_acta)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Informe Conciliación"
    fila = 1

    def celda(col: int, valor: Any, font=_F_NORMAL, moneda=False, borde=False):
        c = ws.cell(row=fila, column=col, value=_texto_seguro(_num(valor)))
        c.font = font
        if moneda:
            c.number_format = FMT_MONEDA
            c.alignment = _DERECHA
        if borde:
            c.border = _BORDE_TOTAL
        return c

    def encabezado_tabla(titulos: List[str]):
        for i, t in enumerate(titulos, 1):
            c = celda(i, t, _F_HEADER)
            c.fill = _FILL
            if t.startswith("Monto"):
                c.alignment = _DERECHA

    # 1. Encabezado
    celda(1, "INFORME DE CONCILIACIÓN BANCARIA", _F_TITULO); fila += 1
    celda(1, f"{d.get('empresa')} | Periodo: {d.get('periodo_nombre')} | Estado del acta: {d.get('estado')}", _F_SUB)
    fila += 2
    for etiqueta, valor in (
        ("Banco:", d.get("banco")),
        ("Número de cuenta:", d.get("numero_cuenta")),
        ("Cuenta contable:", d.get("cuenta_contable")),
        ("Periodo:", f"{d.get('fecha_inicio')} al {d.get('fecha_fin')}"),
        ("Moneda:", "USD ($)"),
    ):
        celda(1, etiqueta, _F_BOLD); celda(2, valor); fila += 1
    fila += 1

    # 2. Extracto bancario
    celda(1, "A. EXTRACTO BANCARIO", _F_SECCION); fila += 1
    encabezado_tabla(["Concepto", "Monto ($)"]); fila += 1
    sin_dato = "Sin dato (cargar extracto con saldos)"
    for concepto, monto, negrita in (
        (f"Saldo inicial según extracto (antes del {d.get('fecha_saldo_inicial') or d.get('fecha_inicio')})",
         d.get("saldo_extracto_inicial"), False),
        ("(+) Créditos del periodo", d.get("total_creditos"), False),
        ("(-) Débitos del periodo", d.get("total_debitos"), False),
        ("Saldo final calculado", d.get("saldo_extracto_calculado"), True),
        (f"Saldo final reportado por el banco al {d.get('fecha_saldo_final') or d.get('fecha_fin')}",
         d.get("saldo_extracto_reportado"), True),
        ("Diferencia calculado vs. reportado (debe ser 0)", d.get("diferencia_extracto"), True),
    ):
        celda(1, concepto, _F_BOLD if negrita else _F_NORMAL, borde=negrita)
        celda(2, monto if monto is not None else sin_dato, _F_BOLD if negrita else _F_NORMAL,
              moneda=monto is not None, borde=negrita)
        fila += 1
    fila += 1

    # 3. Conciliación del movimiento del periodo
    celda(1, "B. CONCILIACIÓN DEL MOVIMIENTO DEL PERIODO", _F_SECCION); fila += 1
    encabezado_tabla(["Concepto", "Monto ($)"]); fila += 1
    nc, nd = _suma(d.get("notas_credito", [])), _suma(d.get("notas_debito", []))
    dt, pt = _suma(d.get("depositos_transito", [])), _suma(d.get("pagos_transito", []))
    no_explicada = Decimal(str(d.get("diferencia_no_explicada") or 0))
    for concepto, monto, negrita in (
        ("Movimiento neto según extracto bancario", d.get("neto_banco"), False),
        ("Movimiento neto según libros (Business Central)", d.get("neto_libros"), False),
        ("Diferencia (extracto - libros)", d.get("diferencia"), True),
        ("Explicada por:", None, False),
        ("   (+) Notas de crédito del banco no registradas en libros", nc, False),
        ("   (+) Notas de débito del banco no registradas en libros", nd, False),
        ("   (-) Depósitos en tránsito (en libros, no en el banco)", -dt, False),
        ("   (-) Cheques y pagos en tránsito (en libros, no en el banco)", -pt, False),
        ("Diferencia no explicada", no_explicada, True),
    ):
        celda(1, concepto, _F_BOLD if negrita else _F_NORMAL, borde=negrita)
        if monto is not None:
            font = _F_ALERTA if (concepto == "Diferencia no explicada" and no_explicada != 0) else (
                _F_BOLD if negrita else _F_NORMAL)
            celda(2, monto, font, moneda=True, borde=negrita)
        fila += 1
    fila += 2

    # 4. Detalle de partidas
    celda(1, "C. DETALLE DE PARTIDAS CONCILIATORIAS", _F_SECCION); fila += 2
    for titulo, clave in (
        ("C.1 Notas de crédito del banco no registradas en libros", "notas_credito"),
        ("C.2 Notas de débito del banco no registradas en libros", "notas_debito"),
        ("C.3 Depósitos en tránsito", "depositos_transito"),
        ("C.4 Cheques y pagos en tránsito", "pagos_transito"),
    ):
        items = d.get(clave, [])
        celda(1, f"{titulo} ({len(items)})", _F_BOLD); fila += 1
        encabezado_tabla(["Fecha", "N° Documento / Referencia", "Concepto", "Monto ($)"]); fila += 1
        for it in items:
            celda(1, it.get("fecha")); celda(2, it.get("numero")); celda(3, it.get("concepto"))
            celda(4, it.get("monto", 0), moneda=True)
            fila += 1
        celda(1, "Total", _F_BOLD, borde=True); celda(4, _suma(items), _F_BOLD, moneda=True, borde=True)
        fila += 2

    reversos = d.get("reversos_libros", [])
    if reversos:
        celda(1, "D. ASIENTOS REVERSADOS EN LIBROS (informativo: se anulan entre sí)", _F_SECCION); fila += 1
        celda(1, f"{len(reversos)} asientos con el mismo N° de documento y fecha cuyos montos suman cero.", _F_SUB); fila += 1
        encabezado_tabla(["Fecha", "N° Documento", "Concepto", "Monto ($)"]); fila += 1
        for it in reversos:
            celda(1, it.get("fecha")); celda(2, it.get("numero")); celda(3, it.get("concepto"))
            celda(4, it.get("monto", 0), moneda=True)
            fila += 1
        celda(1, "Total (debe ser 0)", _F_BOLD, borde=True); celda(4, _suma(reversos), _F_BOLD, moneda=True, borde=True)
        fila += 2

    celda(1, f"Generado por {d.get('generado_por') or 'el sistema'} el {d.get('generado_el')}", _F_SUB)

    for col, ancho in zip("ABCD", (58, 26, 50, 18)):
        ws.column_dimensions[col].width = ancho

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out
