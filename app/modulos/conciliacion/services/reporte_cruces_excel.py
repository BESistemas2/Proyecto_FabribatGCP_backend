"""Excel "Detalle de cruces" (ver ConciliacionService.obtener_datos_reporte_cruces)."""
import io
from datetime import date
from typing import Any, Callable, List, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.modulos.conciliacion.services.reporte_excel import (
    FMT_MONEDA, _F_BOLD, _F_HEADER, _F_NORMAL, _F_SECCION, _F_SUB, _F_TITULO, _FILL, _num, _texto_seguro,
)

FMT_FECHA = "DD/MM/YYYY"
_FRANJA = PatternFill(start_color="EEF3F8", end_color="EEF3F8", fill_type="solid")
_PARCIAL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
_NO_CUADRA = Font(name="Calibri", size=10, bold=True, color="C00000")
_SI_CUADRA = Font(name="Calibri", size=10, color="2E7D32")
_ARRIBA = Alignment(vertical="top", wrap_text=False)
_ARRIBA_AJUSTE = Alignment(vertical="top", wrap_text=True)
_CENTRO = Alignment(horizontal="center", vertical="top")
# Encabezados por origen: banco (azul), Business Central (verde), comparación (gris)
_FILL_BC = PatternFill(start_color="1B6E5A", end_color="1B6E5A", fill_type="solid")
_FILL_CMP = PatternFill(start_color="4D4D4D", end_color="4D4D4D", fill_type="solid")


def _fill_encabezado(titulo: str) -> PatternFill:
    if titulo.startswith("BC ·"):
        return _FILL_BC
    if titulo.startswith("Banco ·"):
        return _FILL
    return _FILL_CMP

# (título, ancho, tipo, clave o función)
Columna = Tuple[str, float, str, Any]


def _valor(fila: dict, clave: Any):
    return clave(fila) if callable(clave) else fila.get(clave)


def _escribir(c, valor: Any, tipo: str):
    if isinstance(valor, date):
        c.value = valor
        c.number_format = FMT_FECHA
    else:
        c.value = _texto_seguro(_num(valor))
    if tipo == "moneda":
        c.number_format = FMT_MONEDA
    c.font = _F_NORMAL
    c.alignment = _ARRIBA_AJUSTE if tipo == "ajuste" else (_CENTRO if tipo in ("centro", "fecha") else _ARRIBA)


def _cuenta(d: dict) -> str:
    banco, numero = str(d.get("banco") or ""), str(d.get("numero_cuenta") or "")
    return banco if numero and numero in banco else f"{banco} {numero}".strip()


def _encabezado(ws: Worksheet, titulo: str, d: dict, subtitulo: str = "") -> int:
    ws.cell(row=1, column=1, value=titulo).font = _F_TITULO
    ws.cell(row=2, column=1, value=(
        f"{d['empresa']} | {_cuenta(d)} (BC: {d['codigo_bc']}) | "
        f"Periodo: {d['periodo_nombre']} | Acta: {d['estado_acta']}")).font = _F_SUB
    if subtitulo:
        ws.cell(row=3, column=1, value=subtitulo).font = _F_SUB
    return 5


def _tabla(ws: Worksheet, fila: int, columnas: Sequence[Columna], filas: List[dict],
           relleno: Callable[[dict], Any] = None, despues: Callable[[Any, dict, int], None] = None,
           fijas: int = 0, por_origen: bool = False) -> int:
    for i, (titulo, ancho, tipo, _) in enumerate(columnas, 1):
        c = ws.cell(row=fila, column=i, value=titulo)
        c.font, c.fill = _F_HEADER, (_fill_encabezado(titulo) if por_origen else _FILL)
        c.alignment = Alignment(horizontal={"moneda": "right", "centro": "center", "fecha": "center"}.get(tipo, "left"),
                                vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[fila].height = 30
    inicio = fila
    for f in filas:
        fila += 1
        fill = relleno(f) if relleno else None
        for i, (_, _, tipo, clave) in enumerate(columnas, 1):
            c = ws.cell(row=fila, column=i)
            _escribir(c, _valor(f, clave), tipo)
            if fill:
                c.fill = fill
            if despues:
                despues(c, f, i)
    if not filas:
        fila += 1
        ws.cell(row=fila, column=1, value="Sin registros.").font = _F_SUB
    ws.freeze_panes = ws.cell(row=inicio + 1, column=fijas + 1)
    ws.auto_filter.ref = f"A{inicio}:{get_column_letter(len(columnas))}{max(fila, inicio + 1)}"
    return fila


# ---------------------------------------------------------------------------

def _hoja_resumen(ws: Worksheet, d: dict):
    ws.title = "Resumen"
    fila = _encabezado(ws, "DETALLE DE CRUCES — RESUMEN", d,
                       f"Generado por {d['generado_por'] or 'usuario'} el {d['generado_el']}. "
                       f"Refleja el estado actual de los cruces; no ejecuta el motor.")
    anchos = (44, 58, 12, 14, 14, 18)
    for i, a in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = a

    ws.cell(row=fila, column=1, value="A. CRUCES POR REGLA").font = _F_SECCION
    fila += 1
    titulos = ("Regla", "Qué valida", "Grupos", "Mov. banco", "Asientos BC", "Monto banco ($)")
    for i, t in enumerate(titulos, 1):
        c = ws.cell(row=fila, column=i, value=t)
        c.font, c.fill = _F_HEADER, _FILL
    tot = [0, 0, 0, 0]
    for r in d["resumen_reglas"]:
        fila += 1
        nombre = f"{r['nombre']}{'  (parcial)' if r['parcial'] else ''}"
        valores = (nombre, r["descripcion"], r["grupos"], r["banco"], r["asientos"], r["monto"])
        for i, v in enumerate(valores, 1):
            c = ws.cell(row=fila, column=i)
            _escribir(c, v, "moneda" if i == 6 else ("ajuste" if i == 2 else "texto"))
            if r["parcial"]:
                c.fill = _PARCIAL
        for k, v in enumerate((r["grupos"], r["banco"], r["asientos"], r["monto"])):
            tot[k] += v
    if not d["resumen_reglas"]:
        fila += 1
        ws.cell(row=fila, column=1, value="Sin cruces en el periodo.").font = _F_SUB
    fila += 1
    ws.cell(row=fila, column=1, value="Total").font = _F_BOLD
    for i, v in enumerate(tot, 3):
        c = ws.cell(row=fila, column=i, value=_num(v))
        c.font = _F_BOLD
        if i == 6:
            c.number_format = FMT_MONEDA

    c = d["control"]
    fila += 2
    ws.cell(row=fila, column=1, value="B. CONTROL DE REGISTROS").font = _F_SECCION
    for etiqueta, valor, moneda in (
        ("Movimientos del banco en el periodo", c["banco_total"], False),
        ("   Conciliados", c["banco_conciliados"], False),
        ("   Pendientes", c["banco_pendientes"], False),
        ("   Monto pendiente del banco ($)", c["monto_pend_banco"], True),
        ("Asientos de BC en el periodo", c["erp_total"], False),
        ("   Cruzados con el banco (de este periodo)", c["erp_cruzados"] - c["erp_cruzados_fuera_periodo"], False),
        ("   Reversos dentro de BC", c["erp_reversos"], False),
        ("   Pendientes", c["erp_pendientes"], False),
        ("   Monto pendiente de BC ($)", c["monto_pend_erp"], True),
        ("Asientos de BC de otro periodo cruzados con este", c["erp_cruzados_fuera_periodo"], False),
    ):
        fila += 1
        ws.cell(row=fila, column=1, value=etiqueta).font = _F_NORMAL
        v = ws.cell(row=fila, column=2, value=_num(valor))
        v.font, v.alignment = _F_BOLD, Alignment(horizontal="left")
        if moneda:
            v.number_format = FMT_MONEDA
    otros = c["banco_total"] - c["banco_conciliados"] - c["banco_pendientes"]
    if otros:
        fila += 1
        ws.cell(row=fila, column=1, value=f"   Otros estados: {otros}").font = _F_SUB

    fila += 2
    ws.cell(row=fila, column=1, value="C. PARÁMETROS DEL MOTOR").font = _F_SECCION
    for etiqueta, valor in d["parametros"]:
        fila += 1
        ws.cell(row=fila, column=1, value=etiqueta).font = _F_NORMAL
        ws.cell(row=fila, column=2, value=valor).font = _F_BOLD

    fila += 2
    ws.cell(row=fila, column=1, value="D. CÓMO LEER EL REPORTE").font = _F_SECCION
    for texto in (
        "Cruzados: una fila por cada pareja banco ↔ asiento de BC. En los cruces 1:N el movimiento del banco "
        "se repite con el mismo N° de grupo: filtre por 'Grupo' para ver el cruce completo.",
        "Las filas en amarillo son cruces parciales (sin referencia que los respalde): conviene revisarlas.",
        "'Días' = fecha de BC menos fecha del banco (positivo: BC registró después).",
        "'Referencia que coincidió' se calcula al generar el reporte con la normalización actual del motor.",
        "Pendientes: la columna 'Observación' indica por qué el motor no encontró la contrapartida.",
    ):
        fila += 1
        cel = ws.cell(row=fila, column=1, value=f"• {texto}")
        cel.font, cel.alignment = _F_NORMAL, Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=6)
        ws.row_dimensions[fila].height = 28


def _hoja_cruzados(ws: Worksheet, d: dict):
    fila = _encabezado(ws, "CRUZADOS — BANCO ↔ BUSINESS CENTRAL", d,
                       "Encabezados: azul = extracto del banco · verde = Business Central · gris = comparación. "
                       "Cada dato del banco va junto a su equivalente de BC.")
    columnas: List[Columna] = [
        ("Grupo", 8, "centro", "grupo"),
        ("Regla", 28, "texto", "regla_nombre"),
        ("Estado", 18, "texto", "estado"),
        # Fechas
        ("Banco · Fecha", 12, "fecha", "banco_fecha"),
        ("BC · Fecha", 12, "fecha", "bc_fecha"),
        ("Días (BC − banco)", 10, "centro", "dias"),
        # Referencias
        ("Banco · Referencia", 18, "texto", "banco_referencia"),
        ("BC · Doc. externo", 18, "texto", "bc_externo"),
        ("BC · N° documento", 18, "texto", "bc_documento"),
        ("Referencia que coincidió", 20, "texto", "referencia"),
        # Montos
        ("Banco · Monto", 14, "moneda", "banco_monto"),
        ("BC · Monto", 14, "moneda", "bc_monto"),
        ("Suma BC del grupo", 14, "moneda", "suma_bc"),
        ("Cuadra", 8, "centro", lambda f: "Sí" if f["cuadra"] else "No"),
        # Detalle
        ("Banco · Concepto", 32, "texto", "banco_concepto"),
        ("BC · Descripción", 34, "texto", "bc_descripcion"),
        ("BC · Entry No.", 11, "centro", "bc_entry"),
        ("Banco · Info adicional", 40, "texto", "banco_info"),
    ]
    col_cuadra = next(i for i, c in enumerate(columnas, 1) if c[0] == "Cuadra")

    def relleno(f):
        if f["estado"] == "CONCILIADO_PARCIAL":
            return _PARCIAL
        return _FRANJA if f["grupo"] % 2 == 0 else None

    def despues(c, f, i):
        if i == col_cuadra:
            c.font = _SI_CUADRA if f["cuadra"] else _NO_CUADRA

    # Grupo, Regla y Estado quedan fijos al desplazarse a la derecha
    _tabla(ws, fila, columnas, d["cruzados"], relleno, despues, fijas=3, por_origen=True)


def _hoja_reversos(ws: Worksheet, d: dict):
    fila = _encabezado(ws, "REVERSOS DENTRO DE BUSINESS CENTRAL", d,
                       "Asientos con el mismo N° de documento y fecha que se anulan entre sí (+X / −X).")
    columnas: List[Columna] = [
        ("N° documento", 20, "texto", "documento"),
        ("Fecha", 12, "fecha", "fecha"),
        ("Entry No.", 11, "centro", "entry"),
        ("Doc. externo", 18, "texto", "externo"),
        ("Descripción", 44, "texto", "descripcion"),
        ("Monto", 14, "moneda", "monto"),
    ]
    docs = []
    for r in d["reversos"]:
        if r["documento"] not in docs:
            docs.append(r["documento"])
    _tabla(ws, fila, columnas, d["reversos"], lambda r: _FRANJA if docs.index(r["documento"]) % 2 else None)


def _hoja_pend_banco(ws: Worksheet, d: dict):
    fila = _encabezado(ws, "PENDIENTES DEL BANCO", d, "Movimientos del extracto sin asiento de BC conciliado.")
    columnas: List[Columna] = [
        ("Fecha", 12, "fecha", "fecha"),
        ("Referencia", 18, "texto", "referencia"),
        ("Concepto", 34, "texto", "concepto"),
        ("Info adicional", 34, "texto", "info"),
        ("Tipo", 10, "texto", "tipo"),
        ("Monto", 14, "moneda", "monto"),
        ("Observación", 70, "ajuste", "observacion"),
    ]
    _tabla(ws, fila, columnas, d["pendientes_banco"])


def _hoja_pend_bc(ws: Worksheet, d: dict):
    fila = _encabezado(ws, "PENDIENTES DE BUSINESS CENTRAL", d, "Asientos de BC sin movimiento del banco conciliado.")
    columnas: List[Columna] = [
        ("Entry No.", 11, "centro", "entry"),
        ("Fecha", 12, "fecha", "fecha"),
        ("N° documento", 18, "texto", "documento"),
        ("Doc. externo", 18, "texto", "externo"),
        ("Descripción", 38, "texto", "descripcion"),
        ("Tipo", 10, "texto", "tipo"),
        ("Monto", 14, "moneda", "monto"),
        ("Observación", 70, "ajuste", "observacion"),
    ]
    _tabla(ws, fila, columnas, d["pendientes_erp"])


def generar_excel_detalle_cruces(d: dict) -> io.BytesIO:
    wb = Workbook()
    _hoja_resumen(wb.active, d)
    _hoja_cruzados(wb.create_sheet("Cruzados"), d)
    _hoja_reversos(wb.create_sheet("Reversos BC"), d)
    _hoja_pend_banco(wb.create_sheet("Pendientes banco"), d)
    _hoja_pend_bc(wb.create_sheet("Pendientes BC"), d)
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = ws.title != "Resumen"
        # Impresión: horizontal, todas las columnas en el ancho de la hoja
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def nombre_archivo_cruces(d: dict) -> str:
    limpio = lambda s: "".join(ch if ch.isalnum() else "_" for ch in str(s)).strip("_")
    return f"Detalle_Cruces_{limpio(d['banco'])}_{limpio(d['codigo_bc'])}_{limpio(d['periodo_nombre'])}.xlsx"
